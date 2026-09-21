"""Excel template generation, parsing/validation and export (SPEC §2.2, DEVIN A3)."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.constants import (
    INTERFACE_COLUMNS,
    INTERFACE_REQUIRED,
    INTERFACE_SHEET,
    SYSTEM_COLUMNS,
    SYSTEM_REQUIRED,
    SYSTEM_SHEET,
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")

# Sample rows written into the template (SPEC §7)
SAMPLE_INTERFACES: list[list[Any]] = [
    [
        "001_SAP_CRM",
        "고객정보연동",
        "SAP-DB",
        "SAP-EAI-CRM(DB)-SAP",
        "SAP",
        "CRM",
        "Real Time",
        "SAP에서 고객정보가 발생하면 CRM으로 전송한다",
    ],
    [
        "002_CRM_SAP",
        "고객클레임정보",
        "DB-SAP",
        "CRM(DB)-EAI-SAP-CRM",
        "CRM",
        "SAP",
        "Batch",
        "CRM에서 클레임 정보가 발생하면 SAP로 변경 데이터를 전송한다.",
    ],
    [
        "003_HR_SAP",
        "고객정보 조회",
        "JSON-SAP",
        "WEB(JSON)-EAI-SAP-WEB",
        "HR",
        "SAP",
        "Real Time",
        "HR시스템에서 고객 정보 필요시 web service를 통해서 SAP를 조회 후 수신한다.",
    ],
]
SAMPLE_SYSTEMS: list[list[Any]] = [
    [
        1,
        "운영",
        "ERP",
        "ERP 시스템",
        "SAP",
        "10.1.1.1",
        None,
        "erp01",
        "erp01",
        "SAP",
        "SAP ERP 시스템",
    ],
    [
        2,
        "운영",
        "DB",
        "고객관리시스템",
        "CRM",
        "11.1.1.1",
        5422,
        "crmuser01",
        "1234",
        "Postgrese",
        "고객관리시스템",
    ],
    [
        3,
        "운영",
        "DB",
        "HR 시스템",
        "HR",
        "12.2.2.2",
        1521,
        "hruser01",
        "1234",
        "MS SQL",
        "HR 시스템",
    ],
]


@dataclass
class RowError:
    """Validation error for one row."""

    row: int
    field: str
    reason: str


@dataclass
class ParsedSheet:
    """Result of parsing one sheet: valid rows (as dicts keyed by DB field) and errors."""

    rows: list[dict[str, Any]] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)


class ExcelFormatError(ValueError):
    """Raised when the workbook does not match the template layout."""


# --------------------------------------------------------------------------- writing


def _write_header(ws: Worksheet, headers: list[str]) -> None:
    """Write a styled header row and auto-size columns."""
    ws.append(headers)
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = max(14, len(header) * 2 + 4)
    ws.freeze_panes = "A2"


def build_workbook(interfaces: list[list[Any]] | None, systems: list[list[Any]] | None) -> bytes:
    """Create the two-sheet workbook with the given data rows and return the xlsx bytes."""
    wb = Workbook()
    ws_if = wb.active
    ws_if.title = INTERFACE_SHEET
    _write_header(ws_if, list(INTERFACE_COLUMNS))
    for row in interfaces or []:
        ws_if.append(row)

    ws_sys = wb.create_sheet(SYSTEM_SHEET)
    _write_header(ws_sys, list(SYSTEM_COLUMNS))
    for row in systems or []:
        ws_sys.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template() -> bytes:
    """Template workbook: headers plus sample rows 2-4."""
    return build_workbook(SAMPLE_INTERFACES, SAMPLE_SYSTEMS)


# --------------------------------------------------------------------------- reading


def _cell_str(value: Any) -> str | None:
    """Normalise a cell value to a stripped string (None if empty)."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _read_sheet(
    content: bytes, sheet_name: str, columns: dict[str, str], required: tuple[str, ...]
) -> ParsedSheet:
    """Read `sheet_name`, validate headers / required cells and return raw row dicts."""
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated exception types
        raise ExcelFormatError("Not a valid .xlsx file") from exc
    if sheet_name not in wb.sheetnames:
        raise ExcelFormatError(f"Sheet '{sheet_name}' not found")
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if header_row is None:
        raise ExcelFormatError(f"Sheet '{sheet_name}' is empty")
    headers = [_cell_str(h) for h in header_row]
    missing = [h for h in columns if h not in headers]
    if missing:
        raise ExcelFormatError(f"Missing columns in '{sheet_name}': {', '.join(missing)}")
    index = {h: headers.index(h) for h in columns}

    parsed = ParsedSheet()
    for row_no, values in enumerate(rows_iter, start=2):
        if values is None or all(_cell_str(v) is None for v in values):
            continue
        record: dict[str, Any] = {}
        row_ok = True
        for header, db_field in columns.items():
            pos = index[header]
            raw = values[pos] if pos < len(values) else None
            text = _cell_str(raw)
            if header in required and text is None:
                parsed.errors.append(RowError(row_no, header, "required value is empty"))
                row_ok = False
            record[db_field] = text
        if row_ok:
            record["_row"] = row_no
            parsed.rows.append(record)
    wb.close()
    return parsed


def parse_interfaces(content: bytes) -> ParsedSheet:
    """Parse the '인터페이스 리스트' sheet."""
    return _read_sheet(content, INTERFACE_SHEET, INTERFACE_COLUMNS, INTERFACE_REQUIRED)


def parse_systems(content: bytes) -> ParsedSheet:
    """Parse the '시스템 연동정보' sheet, converting `Port` to int."""
    parsed = _read_sheet(content, SYSTEM_SHEET, SYSTEM_COLUMNS, SYSTEM_REQUIRED)
    valid: list[dict[str, Any]] = []
    for record in parsed.rows:
        port = record.get("port")
        if port is not None:
            try:
                record["port"] = int(port)
            except ValueError:
                parsed.errors.append(RowError(record["_row"], "Port", f"'{port}' is not a number"))
                continue
        valid.append(record)
    parsed.rows = valid
    return parsed
