"""Tests for /api/upload (template, upload, export, history)."""

from __future__ import annotations

import io
from typing import Any

from httpx import AsyncClient
from openpyxl import Workbook, load_workbook

from app.constants import INTERFACE_COLUMNS, INTERFACE_SHEET, SYSTEM_COLUMNS, SYSTEM_SHEET
from app.services.excel_parser import SAMPLE_INTERFACES, SAMPLE_SYSTEMS, build_workbook


def _upload_files(content: bytes, name: str = "data.xlsx") -> dict[str, Any]:
    """Multipart payload helper."""
    return {"file": (name, content, "application/octet-stream")}


async def test_template_layout(client: AsyncClient) -> None:
    """Template has both sheets, exact headers and sample rows 2-4."""
    resp = await client.get("/api/upload/template")
    assert resp.status_code == 200
    assert "attachment" in resp.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == [INTERFACE_SHEET, SYSTEM_SHEET]
    ws = wb[INTERFACE_SHEET]
    assert [c.value for c in ws[1]] == list(INTERFACE_COLUMNS)
    assert ws.max_row == 4
    ws2 = wb[SYSTEM_SHEET]
    assert [c.value for c in ws2[1]] == list(SYSTEM_COLUMNS)
    assert ws2.max_row == 4


async def test_upload_systems_then_interfaces(client: AsyncClient) -> None:
    """Happy path: systems sheet then interfaces sheet from the template."""
    template = (await client.get("/api/upload/template")).content
    sys_resp = await client.post("/api/upload/systems", files=_upload_files(template))
    assert sys_resp.status_code == 200, sys_resp.text
    assert sys_resp.json()["success_count"] == 3
    assert sys_resp.json()["skipped_count"] == 0

    if_resp = await client.post("/api/upload/interfaces", files=_upload_files(template))
    assert if_resp.json() == {
        "file_name": "data.xlsx",
        "sheet": INTERFACE_SHEET,
        "success_count": 3,
        "skipped_count": 0,
        "errors": [],
    }

    listed = await client.get("/api/interfaces")
    assert listed.json()["total"] == 3
    crm = await client.get("/api/systems", params={"keyword": "CRM"})
    assert crm.json()["items"][0]["port"] == 5422
    detail = await client.get(
        f"/api/systems/{crm.json()['items'][0]['id']}", params={"include_password": True}
    )
    assert detail.json()["password"] == "1234"


async def test_upload_interfaces_validation(client: AsyncClient) -> None:
    """Duplicates, unknown systems and empty required cells are reported per row."""
    template = (await client.get("/api/upload/template")).content
    await client.post("/api/upload/systems", files=_upload_files(template))
    await client.post("/api/upload/interfaces", files=_upload_files(template))

    rows = [
        SAMPLE_INTERFACES[0],  # duplicate
        ["004_X_Y", "x", "SAP-DB", None, "NOPE", "CRM", None, "Batch", None],  # unknown source
        ["005_X_Y", "x", "SAP-DB", None, "SAP", "NOPE", None, "Batch", None],  # unknown target
        ["006_X_Y", None, "SAP-DB", None, "SAP", "CRM", None, "Batch", None],  # missing name
        ["007_OK", "ok", "SAP-DB", None, "SAP", "CRM", None, "Batch", None],  # valid
        [None] * 9,  # blank row ignored
    ]
    resp = await client.post(
        "/api/upload/interfaces", files=_upload_files(build_workbook(rows, None))
    )
    body = resp.json()
    assert body["success_count"] == 1
    assert body["skipped_count"] == 4
    by_row = {e["row"]: e for e in body["errors"]}
    assert by_row[2]["reason"] == "duplicate interface_id"
    assert by_row[3]["field"] == "소스시스템"
    assert by_row[4]["field"] == "타켓시스템"
    assert by_row[5]["field"] == "인터페이스 이름"


async def test_upload_systems_validation(client: AsyncClient) -> None:
    """Blank system code, bad port and duplicates are skipped."""
    rows = [
        SAMPLE_SYSTEMS[0],
        SAMPLE_SYSTEMS[0],  # duplicate within file
        [
            5,
            "운영",
            "FTP",
            "파일관리 시스템",
            None,
            "14.1.1.1",
            443,
            "u",
            "p",
            "AWS",
            None,
        ],  # blank code
        [6, "운영", "DB", "x", "BADPORT", "1.1.1.1", "abc", None, None, None, None],
    ]
    resp = await client.post("/api/upload/systems", files=_upload_files(build_workbook(None, rows)))
    body = resp.json()
    assert body["success_count"] == 1
    assert body["skipped_count"] == 3
    fields = sorted(e["field"] for e in body["errors"])
    assert fields == ["Port", "시스템코드", "시스템코드"]


async def test_upload_bad_files(client: AsyncClient) -> None:
    """Wrong extension, missing sheet and missing column → 400."""
    assert (
        await client.post("/api/upload/interfaces", files=_upload_files(b"abc", "x.csv"))
    ).status_code == 400

    wb = Workbook()
    wb.active.title = "other"
    buf = io.BytesIO()
    wb.save(buf)
    resp = await client.post("/api/upload/interfaces", files=_upload_files(buf.getvalue()))
    assert resp.status_code == 400
    assert INTERFACE_SHEET in resp.json()["detail"]

    wb = Workbook()
    wb.active.title = INTERFACE_SHEET
    wb.active.append(list(INTERFACE_COLUMNS)[:-2])
    buf = io.BytesIO()
    wb.save(buf)
    resp = await client.post("/api/upload/interfaces", files=_upload_files(buf.getvalue()))
    assert resp.status_code == 400
    assert "Missing columns" in resp.json()["detail"]


async def test_export_roundtrip(client: AsyncClient) -> None:
    """Upload template → export → sheets match uploaded data."""
    template = (await client.get("/api/upload/template")).content
    await client.post("/api/upload/systems", files=_upload_files(template))
    await client.post("/api/upload/interfaces", files=_upload_files(template))

    resp = await client.get("/api/upload/export", params={"include_password": True})
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    if_rows = [list(r) for r in wb[INTERFACE_SHEET].iter_rows(min_row=2, values_only=True)]
    assert if_rows == SAMPLE_INTERFACES
    sys_rows = [list(r) for r in wb[SYSTEM_SHEET].iter_rows(min_row=2, values_only=True)]
    assert sys_rows == SAMPLE_SYSTEMS

    masked = await client.get("/api/upload/export")
    wb2 = load_workbook(io.BytesIO(masked.content))
    pw_col = list(SYSTEM_COLUMNS).index("패스워드") + 1
    assert all(
        r[pw_col - 1] is None for r in wb2[SYSTEM_SHEET].iter_rows(min_row=2, values_only=True)
    )


async def test_upload_history(client: AsyncClient) -> None:
    """Every upload is recorded with counts and status."""
    template = (await client.get("/api/upload/template")).content
    await client.post("/api/upload/systems", files=_upload_files(template, "sys.xlsx"))
    await client.post("/api/upload/interfaces", files=_upload_files(template, "if.xlsx"))
    await client.post("/api/upload/interfaces", files=_upload_files(template, "if.xlsx"))

    resp = await client.get("/api/upload/history")
    body = resp.json()
    assert body["total"] == 3
    statuses = [h["status"] for h in body["items"]]
    assert statuses == ["Failed", "Success", "Success"]
    assert body["items"][0]["skipped_count"] == 3
    assert body["items"][2]["sheet"] == "systems"
    assert body["items"][2]["record_count"] == 3
