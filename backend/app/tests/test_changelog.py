"""Automatic change_log recording (SPEC §2.5) and the changelog list / export endpoints."""

from __future__ import annotations

import io
import json

from httpx import AsyncClient
from openpyxl import load_workbook

from app.routers.changelog import CHANGELOG_COLUMNS, CHANGELOG_SHEET
from app.tests.conftest import create_interface, create_system, login


async def test_crud_produces_change_log_rows(client: AsyncClient) -> None:
    sap = await create_system(client)
    crm = await create_system(client, system_code="CRM", system_name="CRM", type="DB")
    iface = await create_interface(client, sap["id"], crm["id"])

    logs = (await client.get("/api/changelog")).json()
    assert logs["total"] == 3
    newest = logs["items"][0]
    assert newest["table_name"] == "interfaces"
    assert newest["action"] == "CREATE"
    assert newest["record_label"] == "001_SAP_CRM"
    assert newest["username"] == "admin"
    snapshot = json.loads(newest["new_value"])
    assert snapshot["interface_id"] == "001_SAP_CRM"
    sys_create = next(
        i for i in logs["items"] if i["table_name"] == "systems" and i["record_id"] == sap["id"]
    )
    assert json.loads(sys_create["new_value"])["password_encrypted"] == "***"

    # update: one row per changed field, password never leaks, untouched fields skipped
    body = {k: v for k, v in sap.items() if k in {
        "category", "type", "system_name", "system_code", "ip", "port",
        "account", "product_name", "description"}}  # fmt: skip
    body.update(system_name="ERP 신규", port=8000, password="changed!")
    assert (await client.put(f"/api/systems/{sap['id']}", json=body)).status_code == 200
    updates = (
        await client.get("/api/changelog", params={"action": "update", "table": "systems"})
    ).json()["items"]
    changed = {u["field_name"]: (u["old_value"], u["new_value"]) for u in updates}
    assert changed == {
        "system_name": ("ERP 시스템", "ERP 신규"),
        "port": (None, "8000"),
        "password_encrypted": ("***", "***"),
    }

    # delete: snapshot kept so the label survives after the record is gone
    if_body = {
        k: iface[k]
        for k in (
            "interface_id",
            "interface_name",
            "integration_type",
            "process",
            "source_system_id",
            "target_system_id",
            "cycle",
            "description",
            "status",
        )
    }
    if_body["status"] = "Inactive"
    assert (await client.put(f"/api/interfaces/{iface['id']}", json=if_body)).status_code == 200
    assert (await client.delete(f"/api/interfaces/{iface['id']}")).status_code == 204
    deleted = (await client.get("/api/changelog", params={"action": "DELETE"})).json()["items"]
    assert len(deleted) == 1
    assert deleted[0]["record_label"] == "001_SAP_CRM"
    assert deleted[0]["new_value"] is None
    # earlier rows of the deleted record keep the label too
    history = (await client.get("/api/changelog", params={"table": "interfaces"})).json()["items"]
    assert [h["action"] for h in history] == ["DELETE", "UPDATE", "CREATE"]
    assert {h["record_label"] for h in history} == {"001_SAP_CRM"}

    # no-op update writes nothing
    before = (await client.get("/api/changelog")).json()["total"]
    body.pop("password")
    assert (await client.put(f"/api/systems/{sap['id']}", json=body)).status_code == 200
    assert (await client.get("/api/changelog")).json()["total"] == before


async def test_changelog_filters_and_export(client: AsyncClient) -> None:
    sap = await create_system(client)
    created = await client.post("/api/users", json={"username": "kim", "password": "kim12345"})
    kim = created.json()
    kim_tokens = await login(client, "kim", "kim12345")
    kim_headers = {"Authorization": f"Bearer {kim_tokens['access_token']}"}
    await create_system(client, system_code="CRM", headers=kim_headers)

    by_user = (await client.get("/api/changelog", params={"user_id": kim["id"]})).json()
    assert by_user["total"] == 1
    assert by_user["items"][0]["record_label"] == "CRM"
    assert by_user["items"][0]["username"] == "kim"

    by_record = (await client.get("/api/changelog", params={"record_id": sap["id"]})).json()
    assert by_record["total"] == 1
    assert (await client.get("/api/changelog", params={"date_from": "2999-01-01"})).json()[
        "total"
    ] == 0
    assert (await client.get("/api/changelog", params={"date_to": "2000-01-01"})).json()[
        "total"
    ] == 0
    assert (await client.get("/api/changelog", params={"table": "users"})).status_code == 422

    resp = await client.get("/api/changelog/export", params={"table": "systems"})
    assert resp.status_code == 200
    assert "dataflowdesk_changelog_" in resp.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(resp.content))[CHANGELOG_SHEET]
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    assert rows[0] == CHANGELOG_COLUMNS
    assert [r[2] for r in rows[1:]] == ["CRM", "SAP"]
    assert [r[7] for r in rows[1:]] == ["kim", "admin"]
    assert all(r[1] == "시스템" and r[3] == "CREATE" for r in rows[1:])
