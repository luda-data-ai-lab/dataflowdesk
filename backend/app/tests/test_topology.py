"""Tests for /api/topology and the `via_system_id` (hub) field."""

from __future__ import annotations

import io
from typing import Any

from httpx import AsyncClient
from openpyxl import load_workbook

from app.constants import INTERFACE_SHEET
from app.tests.conftest import create_interface, create_system


async def _hub_fixture(client: AsyncClient) -> dict[str, Any]:
    """SAP, CRM, HR + IFSYS hub; one routed interface and one direct one."""
    sap = await create_system(client)
    crm = await create_system(client, system_code="CRM", type="DB", system_name="고객관리시스템")
    hr = await create_system(client, system_code="HR", type="DB", system_name="HR 시스템")
    hub = await create_system(
        client, system_code="IFSYS", type="EAI", system_name="인터페이스 시스템"
    )
    routed = await create_interface(client, sap["id"], crm["id"], via_system_id=hub["id"])
    direct = await create_interface(
        client, hr["id"], sap["id"], interface_id="003_HR_SAP", cycle="Batch"
    )
    return {"sap": sap, "crm": crm, "hr": hr, "hub": hub, "routed": routed, "direct": direct}


async def test_via_system_embedded_and_filterable(client: AsyncClient) -> None:
    """`via_system` is returned and usable in list filters."""
    f = await _hub_fixture(client)
    assert f["routed"]["via_system"]["system_code"] == "IFSYS"
    assert f["direct"]["via_system"] is None

    by_via = await client.get("/api/interfaces", params={"via": "IFSYS"})
    assert [i["interface_id"] for i in by_via.json()["items"]] == ["001_SAP_CRM"]
    by_system = await client.get("/api/interfaces", params={"system": "IFSYS"})
    assert by_system.json()["total"] == 1


async def test_via_unknown_system_400(client: AsyncClient) -> None:
    """Unknown hub id is rejected like source/target."""
    sap = await create_system(client)
    crm = await create_system(client, system_code="CRM")
    resp = await client.post(
        "/api/interfaces",
        json={
            "interface_id": "X",
            "interface_name": "x",
            "integration_type": "SAP-DB",
            "source_system_id": sap["id"],
            "target_system_id": crm["id"],
            "via_system_id": 999,
            "cycle": "Batch",
        },
    )
    assert resp.status_code == 400
    assert "via_system_id" in resp.json()["detail"]


async def test_hub_delete_protected(client: AsyncClient) -> None:
    """A system used only as a hub still counts as referenced."""
    f = await _hub_fixture(client)
    resp = await client.delete(f"/api/systems/{f['hub']['id']}")
    assert resp.status_code == 409
    detail = await client.get(f"/api/systems/{f['hub']['id']}")
    assert detail.json()["interface_count"] == 1


async def test_topology_splits_routed_edges(client: AsyncClient) -> None:
    """Routed interface -> two edges through the hub; direct -> one edge."""
    f = await _hub_fixture(client)
    resp = await client.get("/api/topology")
    assert resp.status_code == 200
    body = resp.json()
    assert body["hub_id"] == f["hub"]["id"]
    assert body["hub_code"] == "IFSYS"
    nodes = {n["system_code"]: n for n in body["nodes"]}
    assert nodes["IFSYS"]["is_hub"] is True
    assert nodes["IFSYS"]["interface_count"] == 1
    assert nodes["SAP"]["interface_count"] == 2

    edges = {(e["source_id"], e["target_id"]): e for e in body["edges"]}
    sap, crm, hr, hub = (f[k]["id"] for k in ("sap", "crm", "hr", "hub"))
    assert set(edges) == {(sap, hub), (hub, crm), (hr, sap)}
    assert edges[(sap, hub)]["interfaces"][0]["interface_id"] == "001_SAP_CRM"
    assert edges[(hr, sap)]["interfaces"][0]["interface_id"] == "003_HR_SAP"


async def test_topology_filters_and_missing_hub(client: AsyncClient) -> None:
    """Category/status filters apply; hub_id is null when the hub code is absent."""
    f = await _hub_fixture(client)
    inactive = await client.get("/api/topology", params={"status": "Inactive"})
    assert inactive.json()["edges"] == []
    assert len(inactive.json()["nodes"]) == 4

    other_hub = await client.get("/api/topology", params={"hub": "NOPE"})
    assert other_hub.json()["hub_id"] is None
    assert all(n["is_hub"] is False for n in other_hub.json()["nodes"])

    dev = await client.get("/api/topology", params={"category": "개발"})
    assert dev.json()["nodes"] == []
    assert f["routed"]["id"]  # fixture used


async def test_excel_via_column_roundtrip(client: AsyncClient) -> None:
    """Export writes 경유시스템; upload resolves it or reports unknown code."""
    f = await _hub_fixture(client)
    export = await client.get("/api/upload/export")
    wb = load_workbook(io.BytesIO(export.content))
    ws = wb[INTERFACE_SHEET]
    headers = [c.value for c in ws[1]]
    via_col = headers.index("경유시스템")
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["001_SAP_CRM"][via_col] == "IFSYS"
    assert rows["003_HR_SAP"][via_col] is None

    # re-upload under new ids: one valid hub, one unknown hub
    ws.cell(row=2, column=1, value="101_NEW")
    ws.cell(row=3, column=1, value="102_NEW")
    ws.cell(row=3, column=via_col + 1, value="GHOST")
    buf = io.BytesIO()
    wb.save(buf)
    resp = await client.post(
        "/api/upload/interfaces", files={"file": ("x.xlsx", buf.getvalue(), "application/x")}
    )
    body = resp.json()
    assert body["success_count"] == 1
    assert body["skipped_count"] == 1
    assert body["errors"][0]["field"] == "경유시스템"
    assert f["hub"]["id"]  # fixture used


async def test_upload_without_via_column_still_works(client: AsyncClient) -> None:
    """Sheets produced before the 경유시스템 column existed remain uploadable."""
    template = (await client.get("/api/upload/template")).content
    await client.post("/api/upload/systems", files={"file": ("t.xlsx", template, "x")})
    wb = load_workbook(io.BytesIO(template))
    ws = wb[INTERFACE_SHEET]
    headers = [c.value for c in ws[1]]
    ws.delete_cols(headers.index("경유시스템") + 1)
    buf = io.BytesIO()
    wb.save(buf)
    resp = await client.post(
        "/api/upload/interfaces", files={"file": ("t.xlsx", buf.getvalue(), "x")}
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["success_count"] == 3
