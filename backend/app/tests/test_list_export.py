"""Tests for the filtered list downloads: /api/interfaces/export and /api/systems/export."""

from __future__ import annotations

import io

from httpx import AsyncClient
from openpyxl import load_workbook

from app.constants import INTERFACE_SHEET, SYSTEM_SHEET
from app.routers.interfaces import INTERFACE_LIST_COLUMNS
from app.routers.systems import SYSTEM_LIST_COLUMNS
from app.tests.conftest import create_interface, create_system


def _rows(content: bytes, sheet: str) -> list[tuple]:
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == [sheet]
    return list(wb[sheet].iter_rows(values_only=True))


async def _fixture(client: AsyncClient) -> None:
    sap = await create_system(client)
    crm = await create_system(client, system_code="CRM", type="DB", category="개발")
    await create_interface(client, sap["id"], crm["id"])
    await create_interface(
        client,
        crm["id"],
        sap["id"],
        interface_id="002_CRM_SAP",
        integration_type="DB-SAP",
        status="Inactive",
    )


async def test_export_interfaces_respects_filters(client: AsyncClient) -> None:
    await _fixture(client)

    resp = await client.get("/api/interfaces/export")
    assert resp.status_code == 200
    assert "dataflowdesk_interfaces_" in resp.headers["content-disposition"]
    rows = _rows(resp.content, INTERFACE_SHEET)
    assert rows[0] == INTERFACE_LIST_COLUMNS
    assert [r[0] for r in rows[1:]] == ["001_SAP_CRM", "002_CRM_SAP"]
    header = list(rows[0])
    assert rows[2][header.index("상태")] == "Inactive"
    assert rows[1][header.index("소스시스템명")] == "ERP 시스템"

    filtered = await client.get("/api/interfaces/export", params={"status": "Active"})
    assert [r[0] for r in _rows(filtered.content, INTERFACE_SHEET)[1:]] == ["001_SAP_CRM"]

    by_system = await client.get(
        "/api/interfaces/export", params={"source": "CRM", "sort": "-interface_id"}
    )
    assert [r[0] for r in _rows(by_system.content, INTERFACE_SHEET)[1:]] == ["002_CRM_SAP"]

    bad = await client.get("/api/interfaces/export", params={"sort": "nope"})
    assert bad.status_code == 400


async def test_export_systems_respects_filters_and_masks_password(client: AsyncClient) -> None:
    await _fixture(client)

    resp = await client.get("/api/systems/export")
    assert resp.status_code == 200
    rows = _rows(resp.content, SYSTEM_SHEET)
    assert rows[0] == SYSTEM_LIST_COLUMNS
    assert "비밀번호" not in rows[0]
    header = list(rows[0])
    assert [r[header.index("시스템코드")] for r in rows[1:]] == ["SAP", "CRM"]
    assert [r[header.index("인터페이스 수")] for r in rows[1:]] == [2, 2]
    assert [r[0] for r in rows[1:]] == [1, 2]

    filtered = await client.get("/api/systems/export", params={"category": "개발"})
    frows = _rows(filtered.content, SYSTEM_SHEET)
    assert [r[header.index("시스템코드")] for r in frows[1:]] == ["CRM"]

    bad = await client.get("/api/systems/export", params={"sort": "password"})
    assert bad.status_code == 422
